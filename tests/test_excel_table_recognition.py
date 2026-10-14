"""Parser regressions using saved Excel files, not fabricated worksheet mocks."""

import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

from report_automation_engine import excel_report_generator as engine


@pytest.fixture
def config():
    return engine.load_style_config(engine.DEFAULT_CONFIG_PATH)


def save_workbook(tmp_path, rows, merges=(), sheets=(), formats=()):
    wb = openpyxl.Workbook()
    try:
        ws = wb.active
        ws.title = "Data"
        for row in rows:
            ws.append(row)
        for merged in merges:
            ws.merge_cells(merged)
        for cell, number_format in formats:
            ws[cell].number_format = number_format
        for name in sheets:
            other = wb.create_sheet(name)
            for row in rows:
                other.append(row)
        path = tmp_path / "tables.xlsx"
        wb.save(path)
        return path
    finally:
        wb.close()


def extract(tmp_path, config, rows, **kwargs):
    path = save_workbook(tmp_path, rows, **kwargs)
    return engine.extract_tables_from_excel(path, config)


@pytest.mark.parametrize("label", [
    "사례수", "표본수", "응답수", "빈도", "N", "케이스수",
    "사 례\n수 (N)", "표 본 수\n(명)", "응답수 ( N )", "(N)", "n (명)",
])
def test_case_aliases_and_multiline_suffixes(tmp_path, config, label):
    tables = extract(tmp_path, config, [
        ["<표> 이용 경험"], ["구분", label, "있음", "없음"],
        ["전체", 100, 65, 35],
    ])
    assert len(tables) == 1
    assert tables[0].title == "이용 경험"
    assert tables[0].total_row.case_count == 100
    assert tables[0].total_row.values == {"있음": 65, "없음": 35}


@pytest.mark.parametrize("title", [
    "<표> 이용 경험", "[ 표 ] 이용 경험", "[ 표 1-2 ] 이용 경험",
    "표 1-2. 이용 경험", "Table 3. 이용 경험", "1. 이용 경험",
    "Q3. 이용 경험", "문 3. 이용 경험", "이용 경험",
    "[표 1-2 > 이용 경험]", "[ 표 1-2 > 1. 이용 경험 ]",
])
@pytest.mark.parametrize("shift", [0, 1, 2])
def test_title_markers_plain_and_shifted_columns(tmp_path, config, title, shift):
    prefix = [None] * shift
    tables = extract(tmp_path, config, [
        prefix + [title], prefix + ["구분", "사례수", "있음", "없음"],
        prefix + ["전체", 100, 65, 35],
    ])
    assert len(tables) == 1
    assert tables[0].title == "이용 경험"
    assert tables[0].total_row.values == {"있음": 65, "없음": 35}


def test_closest_title_not_previous_marker_or_banner(tmp_path, config):
    tables = extract(tmp_path, config, [
        ["[표 1] 이전 표"], ["구분", "사례수", "있음"], ["전체", 100, 60],
        [None, "현재 질문"], ["성별"], ["BASE: 전체", 100],
        ["구분", "사례수", "있음"], ["전체", 80, 70],
    ])
    assert [t.title for t in tables] == ["이전 표", "현재 질문"]
    assert [r.label for r in tables[0].rows] == ["전체"]
    assert tables[0].end_row == 3
    assert tables[1].response_columns == ["있음"]


@pytest.mark.parametrize("prefix", [[], [["BASE: 전체", 100], ["출처: 조사", 2025]],
                                   [["구분", "사례수", "있음"]]])
def test_earliest_header_and_no_title_keep_warning(tmp_path, config, prefix):
    tables = extract(tmp_path, config, prefix + [
        ["구분", "사례수", "있음"], ["전체", 100, 65],
    ])
    assert len(tables) == 1
    assert tables[0].title == engine.TITLE_UNKNOWN
    assert tables[0].total_row.values == {"있음": 65}


def test_no_case_distribution_requires_two_actual_rows(tmp_path, config):
    tables = extract(tmp_path, config, [
        [None, "이용 분포"], [None, "항목", None, "비율"],
        [None, "매일", 999, "60%"], [None, "가끔", 888, "40%"],
    ])
    assert len(tables) == 1
    assert tables[0].title == "이용 분포"
    assert tables[0].response_columns == ["비율"]
    assert [r.values for r in tables[0].rows] == [{"비율": 60}, {"비율": 40}]
    assert all(r.case_count is None for r in tables[0].rows)


@pytest.mark.parametrize("rows", [
    [["사례수는 응답자 수를 의미함", "설명"], ["자료", 2025]],
    [["구분", "사례수", "응답"], ["출처: 조사", 2025, 50]],
    [["구분", "사례수", "응답"], [None, 100, 50]],
    [["구분", "사례수", "응답"], ["전체", "약 100명", "문항 50 참조"]],
    [["항목", "비율"], ["매일", 60], ["출처: 조사", 2025]],
    [["구분", "사례수", "응답"]],
])
def test_prose_notes_unlabeled_and_unbacked_headers_rejected(tmp_path, config, rows):
    assert extract(tmp_path, config, rows) == []


def test_blank_and_duplicate_response_columns_preserve_values(tmp_path, config):
    tables = extract(tmp_path, config, [
        ["이용 경험"], ["구분", "사례수", "있음", None, "있음", "있음 (2)", "없음"],
        ["전체", 100, 11, 999, 22, 33, 44],
        ["남성", 40, 51, None, 52, 53, 54],
    ])
    table = tables[0]
    assert len(table.response_columns) == 4
    assert len(set(table.response_columns)) == 4
    assert list(table.total_row.values.values()) == [11, 22, 33, 44]
    assert list(table.rows[1].values.values()) == [51, 52, 53, 54]
    assert table.response_columns[0] == "있음"
    assert table.response_columns[-1] == "없음"


def test_merged_title_and_double_case_header_do_not_split(tmp_path, config):
    tables = extract(tmp_path, config, [
        ["<표 1> 이용 경험"], ["구분", "사례수", "경험", None],
        [None, None, "있음", "없음"], ["전체", 100, 65, 35],
    ], merges=("A1:D1", "A2:A3", "B2:B3", "C2:D2"))
    assert len(tables) == 1
    assert tables[0].title == "이용 경험"
    assert tables[0].start_row == 2
    assert tables[0].response_columns == ["경험_있음", "경험_없음"]
    assert tables[0].total_row.values == {"경험_있음": 65, "경험_없음": 35}
    assert len(tables[0].rows) == 1


def test_explicit_merged_title_shifted_into_response_columns(tmp_path, config):
    tables = extract(tmp_path, config, [
        [None, None, "<표 1> 이용 경험"], ["구분", "사례수", "있음", "없음"],
        ["전체", 100, 65, 35],
    ], merges=("C1:D1",))
    assert tables[0].title == "이용 경험"
    assert tables[0].total_row.values == {"있음": 65, "없음": 35}


def test_config_skipped_title_still_bounds_previous_table(tmp_path, config):
    config.title_rules["skip_title_keywords"].append("보류")
    tables = extract(tmp_path, config, [
        ["이용 경험"], ["구분", "사례수", "있음"], ["전체", 100, 65],
        ["[표 2] 보류 질문"], ["구분", "사례수", "없음"], ["전체", 80, 40],
    ])
    assert [t.title for t in tables] == ["이용 경험", engine.TITLE_UNKNOWN]
    assert tables[0].end_row == 3


@pytest.mark.parametrize("header", [
    ["대상", "지표 A", "지표 B"], ["항목", "참여 인원", "선호도"],
    ["Category", "Custom response", "Other response"],
])
def test_structural_custom_headers_have_qa_not_guessed_counts(tmp_path, config, header):
    path = save_workbook(tmp_path, [
        ["이용 경험"], header, ["전체", 100, 65], ["남성", 40, 70],
    ])
    tables = engine.extract_tables_from_excel(path, config)
    assert len(tables) == 1
    assert tables[0].response_columns == header[1:]
    assert tables[0].total_row.values == {header[1]: 100, header[2]: 65}
    assert all(r.case_count is None for r in tables[0].rows)
    assert "CASE_COLUMN_UNKNOWN" in tables[0].recognition_warnings
    assert "UNIT_UNKNOWN" in tables[0].recognition_warnings
    assert not tables[0].multi_response
    report = engine.render_excel_report(path, config)
    assert "CASE_COLUMN_UNKNOWN" in report
    assert "UNIT_UNKNOWN" in report
    assert "100%" not in report
    assert "65%" not in report


def test_structural_fallback_preserves_explicit_unit_evidence(tmp_path, config):
    tables = extract(tmp_path, config, [
        ["이용 경험"], ["단위: %"], ["대상", "응답 A", "응답 B"],
        ["전체", 65, 35], ["남성", 70, 30],
    ])
    assert len(tables) == 1
    assert tables[0].unit_text == "단위: %"
    assert "CASE_COLUMN_UNKNOWN" in tables[0].recognition_warnings
    assert "UNIT_UNKNOWN" not in tables[0].recognition_warnings


@pytest.mark.parametrize("rows", [
    [["현재 질문"], ["대상", "지표 A", "지표 B"], ["전체", 100, 65]],
    [["전체", 100, 65], ["대상", "지표 A", "지표 B"], ["남성", 40, 70], ["여성", 60, 60]],
    [[None, 2024, 2025], ["전체", 100, 65], ["남성", 40, 70]],
    [["이용 경험"], ["구분", "2024", "2025"], ["전체", 100, 65], ["남성", 40, 70]],
    [[None, "2024년", "2025년"], ["전체", 100, 65], ["남성", 40, 70]],
])
def test_structural_fallback_requires_boundary_two_rows_and_real_fields(tmp_path, config, rows):
    assert extract(tmp_path, config, rows) == []


@pytest.mark.parametrize("gap", [12, 35])
@pytest.mark.parametrize("header", [["구분", "사례수", "있음"], ["대상", "응답 A", "응답 B"]])
def test_long_blank_gaps_do_not_consume_title_or_header_lookahead(tmp_path, config, gap, header):
    tables = extract(tmp_path, config, [
        [None, "표 1", "이용 경험"], *([[]] * gap), header,
        *([[]] * gap), ["전체", 100, 65], *([[]] * gap), ["남성", 40, 70],
    ])
    assert len(tables) == 1
    assert tables[0].title == "이용 경험"
    assert [r.label for r in tables[0].rows] == ["전체", "남성"]


def test_split_title_cells_with_no_gap(tmp_path, config):
    tables = extract(tmp_path, config, [
        [None, "표 1", "만족도"], ["구분", "사례수", "응답"], ["전체", 100, 65],
    ])
    assert tables[0].title == "만족도"
    assert tables[0].response_columns == ["응답"]


def test_alphanumeric_table_id_preserves_nested_question_brackets(tmp_path, config):
    tables = extract(tmp_path, config, [
        ["[표 MDBR-6 > [리스트 기준] 이용 경험]"],
        ["구분", "사례수", "있음"], ["전체", 100, 65],
    ])
    assert tables[0].title == "[리스트 기준] 이용 경험"


def test_unbacked_header_cannot_borrow_data_past_merged_next_title(tmp_path, config):
    tables = extract(tmp_path, config, [
        ["구분", "사례수", "응답"], ["<표 2> 이용 경험"],
        ["구분", "사례수", "있음"], ["전체", 100, 65],
    ], merges=("A2:C2",))
    assert len(tables) == 1
    assert tables[0].start_row == 3
    assert tables[0].title == "이용 경험"
    assert tables[0].total_row.values == {"있음": 65}


@pytest.mark.parametrize("label", ["1. 항목", "Q3. 항목"])
def test_numbered_numeric_body_labels_are_not_question_titles(tmp_path, config, label):
    tables = extract(tmp_path, config, [
        ["이용 경험"], ["구분", "사례수", "있음"], [label, 100, 65],
    ])
    assert len(tables) == 1
    assert tables[0].rows[0].label == label
    assert tables[0].rows[0].values == {"있음": 65}


@pytest.mark.parametrize("raw,number_format,expected", [
    (0.633, "0.0%", 63.3), (63.3, "0.0", 63.3), (0.633, "0.0", 0.633),
    (0.633, '0.0"%"', 0.633), (0.633, r"0.0\%", 0.633),
    (0.633, '0.0"literal %" %', 63.3), (0.633, "0.0_%", 0.633),
    (0.633, "0.0*%", 0.633), ("63.3%", "0.0%", 63.3),
    (1.5, "0.0%", 150), (-0.633, "0.0%", -63.3),
    (0.633, "0.0;0.0%", 0.633), (-0.633, "0.0;0.0%", -63.3),
    (0.633, "0.0%;0.0", 63.3), (-0.633, "0.0%;0.0", -0.633),
])
def test_response_percentage_uses_real_number_format_not_magnitude(tmp_path, config, raw, number_format, expected):
    path = save_workbook(tmp_path, [
        ["이용 경험"], ["구분", "표본수", "있음"], ["전체", 100, raw],
    ], formats=(("C3", number_format),))
    original = path.read_bytes()
    table = engine.extract_tables_from_excel(path, config)[0]
    assert table.total_row.values["있음"] == pytest.approx(expected)
    assert path.read_bytes() == original
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        assert wb["Data"]["C3"].value == raw
        assert wb["Data"]["C3"].number_format == number_format
    finally:
        wb.close()


def test_count_mean_and_score_fields_are_never_automatically_percent_scaled(tmp_path, config):
    config.ranking_rules["exclude_column_keywords"] = []
    table = extract(tmp_path, config, [
        ["이용 경험"], ["구분", "사례수", "있음", "평균", "5점 점수"],
        ["전체", 0.633, 0.633, 0.633, 0.633],
    ], formats=tuple((f"{col}3", "0.0%") for col in "BCDE"))[0]
    assert table.total_row.case_count == 0.633
    assert table.total_row.values == {"있음": pytest.approx(63.3), "평균": 0.633, "5점 점수": 0.633}


def test_genuine_percentage_format_produces_63_3_percent_narrative(tmp_path, config):
    path = save_workbook(tmp_path, [
        ["이용 경험"], ["구분", "표본수", "있음"], ["전체", 100, 0.633],
    ], formats=(("C3", "0.0%"),))
    report = engine.render_excel_report(path, config)
    assert "63.3%" in report
    assert "0.6%" not in report


def test_structural_fallback_accepts_explicit_percentage_format_as_unit_evidence(tmp_path, config):
    tables = extract(tmp_path, config, [
        ["이용 경험"], ["대상", "응답 A"], ["전체", 0.633], ["남성", 0.7],
    ], formats=(("B3", "0.0%"), ("B4", "0.0%")))
    assert tables[0].total_row.values["응답 A"] == pytest.approx(63.3)
    assert tables[0].recognition_warnings == ("CASE_COLUMN_UNKNOWN",)


def test_numeric_year_labels_in_declared_stub_keep_blank_stub_rejected(tmp_path, config):
    tables = extract(tmp_path, config, [
        ["연도", "N", "응답"], [2023, 100, 60], [2024, 120, 70], [None, 140, 80],
    ])
    assert len(tables) == 1
    assert [r.label for r in tables[0].rows] == ["2023", "2024"]
    assert [r.case_count for r in tables[0].rows] == [100, 120]
    assert [r.values for r in tables[0].rows] == [{"응답": 60}, {"응답": 70}]


def test_unit_header_row_is_not_data_or_title(tmp_path, config):
    tables = extract(tmp_path, config, [
        ["이용 경험"], ["구분", "사례수", "있음", "없음"],
        [None, None, "%", "%"], ["전체", "(100)", 65, 35],
    ])
    assert len(tables) == 1
    assert tables[0].response_columns == ["있음", "없음"]
    assert len(tables[0].rows) == 1


def test_base_merged_through_case_and_percent_header_tiers(tmp_path, config):
    tables = extract(tmp_path, config, [
        ["<표 1> 이용 경험"], ["BASE: 전체", None, None, "경험", None],
        [None, None, "사례수", "있음", "없음"],
        [None, None, None, "%", "%"], ["■ 전체", None, "(100)", 65, 35],
    ], merges=("A1:E1", "A2:B4", "C3:C4", "D2:E2", "A5:B5"))
    assert len(tables) == 1
    assert tables[0].title == "이용 경험"
    assert tables[0].start_row == 3
    assert tables[0].response_columns == ["경험_있음", "경험_없음"]
    assert tables[0].total_row.values == {"경험_있음": 65, "경험_없음": 35}
    assert len(tables[0].rows) == 1


def test_configured_case_takes_precedence_over_earlier_alias(tmp_path, config):
    config.table_rules["case_column_name"] = "응답 인원"
    table = extract(tmp_path, config, [
        ["이용 경험"], ["구분", "N", "응답 인원", "있음"],
        ["전체", 999, 100, 65],
    ])[0]
    assert table.total_row.case_count == 100
    assert table.total_row.values == {"있음": 65}


def test_blank_separator_resumes_and_notes_never_become_numeric_rows(tmp_path, config):
    tables = extract(tmp_path, config, [
        ["이용 경험"], ["구분", "사례수", "있음"], ["전체", 100, 65],
        [], [], [], [], [], ["남성", 40, 70],
        ["출처: 조사", 2025, 99], ["주: 반올림", 1, 99], ["Note: sample", 2, 99],
        ["사례수 설명", None, None],
        ["다음 질문"], ["구분", "N", "없음"], ["전체", 80, 40],
    ])
    assert [t.title for t in tables] == ["이용 경험", "다음 질문"]
    assert [r.label for r in tables[0].rows] == ["전체", "남성"]
    assert tables[1].total_row.values == {"없음": 40}


def test_adjacent_no_title_does_not_inherit_previous_table(tmp_path, config):
    tables = extract(tmp_path, config, [
        ["[표 1] 이용 경험"], ["구분", "사례수", "있음"], ["전체", 100, 65],
        ["구분", "사례수", "없음"], ["전체", 80, 40],
    ])
    assert [t.title for t in tables] == ["이용 경험", engine.TITLE_UNKNOWN]
    assert tables[0].end_row == 3
    assert tables[1].response_columns == ["없음"]


@pytest.mark.parametrize("total", ["전체", "■ 전 체", "● 합계", "- 총계", "• Total", "TOTAL"])
def test_total_aliases_and_optional_bullets(tmp_path, config, total):
    table = extract(tmp_path, config, [
        ["이용 경험"], ["구분", "사례수", "있음"], [total, 100, 65],
    ])[0]
    assert table.total_row is not None
    assert table.total_row.case_count == 100


def test_configured_case_total_title_and_ranking_rules_still_apply(tmp_path, config):
    config.table_rules["case_column_name"] = "응답 인원"
    config.table_rules["total_label"] = "모두"
    config.title_rules["strip_prefixes"].append("질문: ")
    config.ranking_rules["exclude_column_keywords"] = ["제외"]
    config.ranking_rules["multi_response_threshold"] = 120
    tables = extract(tmp_path, config, [
        ["질문: 이용 경험"], ["구분", "응답\n인원 (명)", "있음", "제외"],
        ["모두", 100, 110, 10],
    ])
    assert len(tables) == 1
    assert tables[0].title == "이용 경험"
    assert tables[0].total_row.case_count == 100
    assert tables[0].ranked_columns == ["있음"]
    assert not tables[0].multi_response


def test_sheet_filter_index_skip_and_max_tables(tmp_path, config):
    path = save_workbook(tmp_path, [
        ["이용 경험"], ["구분", "사례수", "있음"], ["전체", 100, 65],
    ], sheets=("INDEX", "Other"))
    assert [t.sheet_name for t in engine.extract_tables_from_excel(path, config)] == ["Data", "Other"]
    assert [t.sheet_name for t in engine.extract_tables_from_excel(path, config, sheet_name="Other")] == ["Other"]
    assert len(engine.extract_tables_from_excel(path, config, max_tables=1)) == 1


@pytest.mark.parametrize("header", [["구분", "표본수", "있음"], ["row", "col", "value"]])
def test_raw_rerun_excludes_generated_report_sheets(tmp_path, config, header):
    generated = ("보고서_표", "보고서_분석문", "_ReportCells", "_ReportAnalysis", "_report_preview")
    path = save_workbook(tmp_path, [
        ["이용 경험"], header, ["전체", 100, 65], ["남성", 40, 70],
    ], sheets=generated)
    original = path.read_bytes()
    for _ in range(2):
        tables = engine.extract_tables_from_excel(path, config)
        assert [t.sheet_name for t in tables] == ["Data"]
        assert engine.extract_tables_from_excel(path, config, sheet_name="_ReportCells") == []
    assert path.read_bytes() == original


@pytest.mark.parametrize("extra_args", [["--raw-tables"], []])
def test_cli_import_and_raw_parsing_without_tkinter(tmp_path, extra_args):
    path = save_workbook(tmp_path, [
        ["이용 경험"], ["구분", "표본수", "있음"], ["전체", 100, 65],
    ])
    output = tmp_path / "draft.txt"
    script = """
import builtins
import sys
original_import = builtins.__import__
def no_tk(name, *args, **kwargs):
    if name == 'tkinter' or name.startswith('tkinter.') or name == '_tkinter':
        raise ImportError('Tk is intentionally absent from the standalone payload')
    return original_import(name, *args, **kwargs)
builtins.__import__ = no_tk
from report_automation_engine.excel_report_generator import run_cli
raise SystemExit(run_cli(sys.argv[1:]))
"""
    result = subprocess.run(
        [sys.executable, "-B", "-X", "utf8", "-c", script, "--excel", str(path),
         "--output", str(output), *extra_args],
        cwd=Path(engine.__file__).resolve().parents[1],
        capture_output=True, text=True, encoding="utf-8", timeout=30,
    )
    assert result.returncode == 0, result.stderr
    assert "이용 경험" in output.read_text(encoding="utf-8")
    assert "65%" in output.read_text(encoding="utf-8")


@pytest.mark.parametrize("mode", ["normal", "limit", "error", "fallback", "generated"])
def test_real_loaded_workbooks_close_on_all_exit_paths(tmp_path, config, monkeypatch, mode):
    path = save_workbook(tmp_path, [
        ["이용 경험"], ["구분", "사례수", "있음"], ["전체", 100, 65],
    ], sheets=("보고서_분석문",) if mode == "generated" else ())
    real_load = openpyxl.load_workbook
    closed = []

    def tracked_load(*args, **kwargs):
        wb = real_load(*args, **kwargs)
        real_close = wb.close

        def tracked_close():
            closed.append(wb)
            real_close()

        monkeypatch.setattr(wb, "close", tracked_close)
        return wb

    monkeypatch.setattr(engine.openpyxl, "load_workbook", tracked_load)
    if mode == "error":
        config.ranking_rules["multi_response_threshold"] = "invalid"
        with pytest.raises(ValueError):
            engine.extract_tables_from_excel(path, config)
    elif mode in {"fallback", "generated"}:
        result = engine.render_from_generated_output_sheets(path)
        assert (result is None) == (mode == "fallback")
    else:
        assert len(engine.extract_tables_from_excel(path, config, max_tables=1 if mode == "limit" else None)) == 1
    assert len(closed) == 1
