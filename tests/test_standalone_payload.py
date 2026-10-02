import csv
import importlib.util
import json
import os
import subprocess
import tempfile
import unittest
from unittest.mock import patch
import zipfile
from pathlib import Path, PurePosixPath


SCRIPT = Path(__file__).resolve().parents[1] / "report_automation_launcher/scripts/build_standalone_payload.py"


class StandalonePayloadTests(unittest.TestCase):
    def setUp(self):
        self.assertTrue(SCRIPT.is_file(), "standalone payload builder must exist")
        spec = importlib.util.spec_from_file_location("standalone_payload", SCRIPT)
        self.builder = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(self.builder)
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.repo = self.root / "repo"
        self.runtime = self.root / "python"
        self.site = self.root / "site-packages"
        for name in ["python.exe", "python312.dll", "python3.dll", "vcruntime140.dll", "LICENSE.txt",
                     "Lib/json/__init__.py", "Lib/encodings/__init__.py", "DLLs/_ssl.pyd",
                     "DLLs/_tkinter.pyd", "DLLs/tcl86t.dll", "DLLs/tk86t.dll",
                     "tcl/tcl8.6/init.tcl", "tcl/tcl8.6/license.terms"]:
            self.put(self.runtime / name)
        for name in ["LICENSE", "THIRD_PARTY_NOTICES.md", "report_automation_engine/__init__.py",
                     "report_automation_engine/tool.py", "report_automation_engine/config/default_style_schema.json",
                     "report_automation_engine/config/default_hwp_style_config.json",
                     "report_automation_addin/dev/ReportAutomationAddin_dev.xlam",
                     "templates/report_template_basic.pptx", "templates/chart_review_template_basic.pptx", "docs/template_authoring_guide.md",
                     "report_automation_launcher/licenses/openpyxl-1.0-LICENCE.rst",
                     "report_automation_launcher/licenses/et_xmlfile-1.0-LICENCE.rst",
                     "report_automation_launcher/licenses/et_xmlfile-1.0-LICENCE.python"]:
            self.put(self.repo / name)

    def put(self, path, content="fixture"):
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(content, encoding="utf-8")

    def distribution(self, name, files, requirements=()):
        info = self.site / (name.replace("-", "_") + "-1.0.dist-info")
        self.put(info / "METADATA", "Metadata-Version: 2.1\nName: " + name + "\nVersion: 1.0\n" +
                 "".join("Requires-Dist: " + requirement + "\n" for requirement in requirements))
        self.put(info / "licenses/LICENSE")
        for filename in files:
            if not filename.startswith("../"):
                self.put(self.site / filename)
        records = list(files) + [str(p.relative_to(self.site).as_posix()) for p in info.rglob("*") if p.is_file()]
        with (info / "RECORD").open("w", newline="", encoding="utf-8") as stream:
            csv.writer(stream).writerows((p, "", "") for p in records + [info.name + "/RECORD"])

    def dependencies(self):
        self.distribution("openpyxl", ["openpyxl/__init__.py", "openpyxl/tests/private.py"], ["et-xmlfile>=1"])
        self.distribution("python-pptx", ["pptx/__init__.py", "pptx/templates/default.pptx"],
                          ["Pillow>=1", "lxml>=1", "XlsxWriter>=1", "typing-extensions>=1"])
        self.distribution("Pillow", ["PIL/__init__.py", "PIL/_imaging.pyd"], ["uninstalled; extra == 'tests'"])
        self.distribution("pywin32", ["pywin32.pth", "win32/lib/pywin32_bootstrap.py", "win32/lib/pywintypes.py",
                                     "pythoncom.py", "win32com/__init__.py", "win32/win32api.pyd",
                                     "pywin32_system32/pythoncom312.dll", "pywin32_system32/pywintypes312.dll",
                                     "../../Scripts/pywin32_postinstall.py"])
        for name, filename in [("et-xmlfile", "et_xmlfile/__init__.py"), ("lxml", "lxml/etree.pyd"),
                               ("XlsxWriter", "xlsxwriter/__init__.py"), ("typing-extensions", "typing_extensions.py")]:
            self.distribution(name, [filename])
        self.distribution("unrelated", ["unrelated/private.py"])
        return self.builder.required_distributions([self.site])

    def test_zip_allowlist_manifest_and_dependency_licenses(self):
        distributions = self.dependencies()
        for name in ["outputs/private.py", ".venv/private.py", ".env", "research/private.xlsx",
                     "report_automation_engine/outputs/private.py", "report_automation_engine/tests/private.py",
                     "report_automation_engine/__pycache__/tool.pyc", "report_automation_engine/config/private.txt",
                     "report_automation_engine/config/private.json",
                     "templates/outputs/private.pptx", "templates/private.pptx", "templates/~$private.pptx", "templates/.env"]:
            self.put(self.repo / name)
        for name in ["Lib/site-packages/private.py", "Lib/test/private.py", "Lib/json/__pycache__/private.pyc"]:
            self.put(self.runtime / name)
        files = self.builder.collect_payload_files(self.repo, self.runtime, distributions)
        output = self.root / "payload.zip"
        self.builder.write_payload(output, files, "3.12.14", distributions)
        with zipfile.ZipFile(output) as archive:
            names = archive.namelist()
            manifest = json.loads(archive.read("bundle_manifest.json"))
            self.assertEqual(manifest["files"], sorted(set(names) - {"bundle_manifest.json"}))
            self.assertEqual(manifest["python_version"], "3.12.14")
            self.assertEqual(manifest["architecture"], "x64")
            for generated in (self.builder.RUNTIME_PTH, "bundle_manifest.json"):
                self.assertEqual(archive.getinfo(generated).date_time, (1980, 1, 1, 0, 0, 0))
            self.assertEqual(set(manifest["dependencies"]),
                             {"openpyxl", "python-pptx", "pillow", "pywin32", "et-xmlfile", "lxml", "xlsxwriter", "typing-extensions"})
            for name in ["runtime/python.exe", "runtime/python312.dll", "runtime/LICENSE.txt", "runtime/DLLs/_ssl.pyd",
                         "runtime/Lib/site-packages/pptx/templates/default.pptx",
                         "runtime/Lib/site-packages/pywin32.pth", "runtime/Lib/site-packages/win32/lib/pywin32_bootstrap.py",
                         "runtime/Lib/site-packages/pywin32_system32/pythoncom312.dll",
                         "runtime/Lib/site-packages/pywin32_system32/pywintypes312.dll",
                         "LICENSE", "THIRD_PARTY_NOTICES.md", "licenses/openpyxl-1.0-LICENCE.rst",
                         "licenses/et_xmlfile-1.0-LICENCE.python", "templates/report_template_basic.pptx",
                         "report_automation_addin/dev/ReportAutomationAddin_dev.xlam", "docs/template_authoring_guide.md",
                         "report_automation_engine/config/default_hwp_style_config.json"]:
                self.assertIn(name, names)
            for distribution in distributions.values():
                info = distribution.metadata["Name"].replace("-", "_") + "-1.0.dist-info"
                self.assertIn("runtime/Lib/site-packages/" + info + "/licenses/LICENSE", names)
            self.assertFalse(any("private" in name or "unrelated" in name for name in names))
            self.assertFalse(any(name.startswith("runtime/tcl/") or "_tkinter" in name or "tcl86" in name or "tk86" in name for name in names))
            self.assertFalse(any(PurePosixPath(name).is_absolute() or ".." in PurePosixPath(name).parts or "\\" in name for name in names))
            self.assertIn(b"add_dll_directory", archive.read("runtime/Lib/site-packages/report_automation_runtime.pth"))

    def test_external_notices_are_required_for_exact_dependency_version(self):
        distributions = self.dependencies()
        files = self.builder.collect_payload_files(self.repo, self.runtime, distributions)
        self.assertTrue(hasattr(self.builder, "validate_licenses"), "packaged license validation must exist")
        for info in ("openpyxl-1.0.dist-info", "et_xmlfile-1.0.dist-info"):
            files.pop("runtime/Lib/site-packages/" + info + "/licenses/LICENSE")
        self.builder.validate_licenses(files, distributions)
        files.pop("licenses/openpyxl-1.0-LICENCE.rst")
        with self.assertRaises(ValueError):
            self.builder.validate_licenses(files, distributions)

    def test_missing_required_asset_or_dependency_fails(self):
        distributions = self.dependencies()
        (self.repo / "LICENSE").unlink()
        with self.assertRaises(FileNotFoundError):
            self.builder.collect_payload_files(self.repo, self.runtime, distributions)
        with self.assertRaises(ValueError):
            self.builder.required_distributions([self.root / "missing-site"])

    def test_template_factory_works_without_powershell(self):
        from report_automation_engine.template_factory import create_template
        from pptx import Presentation
        output = self.root / "report template.pptx"
        with patch("report_automation_engine.template_factory.subprocess.run", side_effect=FileNotFoundError("PowerShell unavailable")):
            create_template("pptx_report", output)
        slides = Presentation(output).slides
        self.assertEqual(len(slides), 3)
        self.assertTrue(any("{{CHART}}" in shape.text for slide in slides for shape in slide.shapes if shape.has_text_frame))

    def test_dependency_record_cannot_escape(self):
        self.dependencies()
        self.distribution("openpyxl", ["../../private.py"])
        with self.assertRaises(ValueError):
            self.builder.collect_payload_files(self.repo, self.runtime, self.builder.required_distributions([self.site]))

    def test_missing_recorded_library_fails(self):
        distributions = self.dependencies()
        (self.site / "PIL/_imaging.pyd").unlink()
        with self.assertRaises(FileNotFoundError):
            self.builder.collect_payload_files(self.repo, self.runtime, distributions)

    def test_incompatible_dependency_version_fails(self):
        self.dependencies()
        self.distribution("openpyxl", ["openpyxl/__init__.py"], ["et-xmlfile>=2"])
        with self.assertRaises(ValueError):
            self.builder.required_distributions([self.site])

    def test_symlink_is_not_packaged(self):
        distributions = self.dependencies()
        target = self.repo / "report_automation_engine/linked.py"
        try:
            target.symlink_to(self.repo / "LICENSE")
        except OSError:
            self.skipTest("Windows symlink privilege unavailable")
        with self.assertRaises(ValueError):
            self.builder.collect_payload_files(self.repo, self.runtime, distributions)

    @unittest.skipUnless(os.name == "nt", "Windows native compiler check")
    def test_build_publishes_fixed_name_atomically_and_preserves_previous_on_failure(self):
        source = self.root / "src/Launcher.cs"
        self.put(source, 'static class Launcher { static void Main() { System.Console.WriteLine("first"); } }')
        self.put(source.with_name("StandaloneBundle.cs"), "internal static class StandaloneBundle {}")
        output = self.root / "bin"
        command = ["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File",
                   str(SCRIPT.with_name("build_report_automation_launcher.ps1")), "-DevelopmentOnly",
                   "-SourcePath", str(source), "-OutputDir", str(output)]

        def build():
            return subprocess.run(command, cwd=self.root, capture_output=True, text=True, timeout=60)

        result = build()
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        exe = output / "ReportAutomationLauncher.exe"
        first = exe.read_bytes()
        self.put(source, 'static class Launcher { static void Main() { System.Console.WriteLine("second"); } }')
        # An open executable without FILE_SHARE_DELETE must make File.Replace fail, not erase the old EXE.
        with exe.open("rb"):
            result = build()
            self.assertNotEqual(result.returncode, 0)
            self.assertEqual(exe.read_bytes(), first)
        self.assertEqual([p.name for p in output.iterdir()], ["ReportAutomationLauncher.exe"])
        result = build()
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertNotEqual(exe.read_bytes(), first)
        second = exe.read_bytes()
        self.put(source, "invalid C# source")
        self.assertNotEqual(build().returncode, 0)
        self.assertEqual(exe.read_bytes(), second)
        self.assertEqual([p.name for p in output.iterdir()], ["ReportAutomationLauncher.exe"])


if __name__ == "__main__":
    unittest.main()
